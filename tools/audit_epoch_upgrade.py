#!/usr/bin/env python3
"""Compare current project semantics with the pre-upgrade freeze snapshot."""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent / "work" / "real_epoch"
BASELINE = WORK / "baseline"
ALLOWED_MODELICA = {
    "Foundation/Interfaces/EnvironmentPort.mo",
    "Foundation/Models/EphemerisEnvironmentReader.mo",
    "Foundation/Models/OrbitalEnvironmentCore.mo",
    "Scenarios/DefaultScenario.mo",
    "Scenarios/GeneratedScenario.mo",
    "Scenarios/OrbitConfig.mo",
    "Simulation/CompleteMission.mo",
    "Systems/Four_systems/MechanicsOverall.mo",
}
FROZEN_SHEETS = ("Meta", "GroundStations", "ImagingTargets", "InitialConditions", "Components")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def sheet_hash(sheet: Any) -> str:
    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cells.append({"cell": cell.coordinate, "value": serializable(cell.value), "dataType": cell.data_type})
    payload = json.dumps(cells, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def main() -> int:
    before = json.loads((BASELINE / "modelica_hashes_before.json").read_text(encoding="utf-8"))
    current = {
        str(path.relative_to(ROOT)).replace("\\", "/"): digest(path)
        for path in sorted(ROOT.rglob("*.mo"))
    }
    modified = sorted(name for name in set(before) & set(current) if before[name] != current[name])
    added = sorted(set(current) - set(before))
    deleted = sorted(set(before) - set(current))
    unexpected = sorted((set(modified) | set(added) | set(deleted)) - ALLOWED_MODELICA)

    workbook_before = json.loads((BASELINE / "workbook_semantics_before.json").read_text(encoding="utf-8"))
    workbook = load_workbook(ROOT / "DesignConfig.xlsx", data_only=False, read_only=False)
    frozen = {}
    for name in FROZEN_SHEETS:
        actual = sheet_hash(workbook[name])
        expected = workbook_before["sheets"][name]["semanticSHA256"]
        frozen[name] = {"unchanged": actual == expected, "beforeSHA256": expected, "afterSHA256": actual}

    passed = not unexpected and not deleted and all(item["unchanged"] for item in frozen.values())
    report = {
        "status": "PASS" if passed else "FAIL",
        "modelica": {"modified": modified, "added": added, "deleted": deleted, "unexpected": unexpected},
        "workbookFrozenSheets": frozen,
        "frozenHardwareAndNonOrbitConditionsUnchanged": all(item["unchanged"] for item in frozen.values()),
    }
    (WORK / "freeze_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# 日期与轨道环境升级冻结审计", "", f"- 状态：{report['status']}", "", "## Modelica范围", ""]
    lines.extend(f"- 修改：`{name}`" for name in modified)
    lines.extend(f"- 新增：`{name}`" for name in added)
    if not unexpected:
        lines.append("- 非目标Modelica文件：未变化")
    lines.extend(["", "## Excel冻结区", ""])
    lines.extend(f"- {name}：{'UNCHANGED' if item['unchanged'] else 'CHANGED'}" for name, item in frozen.items())
    (WORK / "freeze_audit.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
