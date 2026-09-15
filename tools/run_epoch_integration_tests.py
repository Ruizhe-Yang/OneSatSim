#!/usr/bin/env python3
"""Run bounded integration tests for scenario parsing and environment generation."""

from __future__ import annotations

import json
import sys
import tempfile
from copy import copy
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent
WORK = ROOT.parent / "work" / "real_epoch"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from parameter_schema import ValidationCollector  # noqa: E402
from prepare_ephemeris import (  # noqa: E402
    EphemerisPreparationError,
    ResourceSettings,
    add_physical_seconds,
    parse_absolute_time,
    prepare_environment,
)
from update_nissa_config import parse_orbit, parse_simulation  # noqa: E402


def set_parameter(workbook: Any, sheet: str, name: str, value: Any, unit: str | None = None) -> None:
    ws = workbook[sheet]
    headers = {str(ws.cell(4, column).value): column for column in range(1, ws.max_column + 1)}
    for row in range(5, ws.max_row + 1):
        if ws.cell(row, headers["Parameter"]).value == name:
            ws.cell(row, headers["Value"]).value = value
            if unit is not None:
                ws.cell(row, headers["Unit"]).value = unit
            return
    raise KeyError(f"{sheet}.{name}")


def check(name: str, condition: bool, details: dict[str, Any], checks: list[dict[str, Any]]) -> None:
    checks.append({"name": name, "status": "PASS" if condition else "FAIL", **details})
    if not condition:
        raise AssertionError(f"{name}: {details}")


def metadata_summary(metadata: dict[str, Any]) -> dict[str, Any]:
    validation = metadata["interpolationValidation"]
    return {
        "startUTC": metadata["startUTC"],
        "endUTC": metadata["endUTC"],
        "durationSeconds": metadata["durationSeconds"],
        "finalSampleInterval": metadata["finalSampleInterval"],
        "tableRows": metadata["tableRows"],
        "tableSHA256": metadata["tableSHA256"],
        "propagationConvergenceMaxPosition_m": metadata["propagationConvergenceMaxPosition_m"],
        "maxPositionError_m": validation["maxPositionError_m"],
        "maxVelocityError_m_s": validation["maxVelocityError_m_s"],
        "maxEarthRotationAngleError_arcsec": validation["maxEarthRotationAngleError_arcsec"],
        "maxEarthRotationOrthogonalityError": validation["maxEarthRotationOrthogonalityError"],
        "maxEarthRotationDeterminantError": validation["maxEarthRotationDeterminantError"],
        "eopStatus": metadata["iers"]["label"],
    }


def main() -> int:
    WORK.mkdir(parents=True, exist_ok=True)
    checks: list[dict[str, Any]] = []
    workbook_path = ROOT / "DesignConfig.xlsx"
    default_scenario = (ROOT / "Scenarios" / "DefaultScenario.mo").read_text(encoding="utf-8")

    # Simulation parsing: the legacy Start/Stop rows are derived, not independent inputs.
    for hours in (6, 24, 48):
        workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        set_parameter(workbook, "Simulation", "Duration", hours, "h")
        collector = ValidationCollector()
        simulation, _ = parse_simulation(workbook, collector)
        check(
            f"parse duration {hours} h",
            not collector.errors and simulation["StartTime"] == 0.0 and simulation["StopTime"] == hours * 3600.0,
            {"stopTime_s": simulation["StopTime"], "errors": [item.render() for item in collector.errors]},
            checks,
        )
    for interval, expected in ((1, True), (2, True), (0.5, False)):
        workbook = load_workbook(workbook_path, data_only=True, read_only=False)
        set_parameter(workbook, "Simulation", "OutputInterval", interval, "s")
        collector = ValidationCollector()
        parse_simulation(workbook, collector)
        check(
            f"output interval {interval} s",
            (not collector.errors) is expected,
            {"accepted": not collector.errors, "expected": expected},
            checks,
        )

    # A Cartesian input with any missing velocity component must be rejected.
    workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    set_parameter(workbook, "Orbit", "VZ0", None, "m/s")
    simulation_collector = ValidationCollector()
    simulation, _ = parse_simulation(workbook, simulation_collector)
    orbit_collector = ValidationCollector()
    parse_orbit(workbook, orbit_collector, default_scenario, simulation["StartTimeUTC"], simulation["TimeZone"])
    check(
        "missing Cartesian velocity is rejected",
        bool(orbit_collector.errors),
        {"errors": [item.render() for item in orbit_collector.errors]},
        checks,
    )

    current_metadata = json.loads((ROOT / "Resources" / "Data" / "GeneratedEphemeris" / "environment_metadata.json").read_text(encoding="utf-8"))
    check(
        "current 24 h generated environment",
        current_metadata["durationSeconds"] == 86400.0
        and current_metadata["interpolationValidation"]["maxPositionError_m"] <= 10.0
        and current_metadata["interpolationValidation"]["maxVelocityError_m_s"] <= 0.02
        and current_metadata["interpolationValidation"]["maxEarthRotationAngleError_arcsec"] <= 1.0,
        metadata_summary(current_metadata),
        checks,
    )

    resources = ResourceSettings(
        spk_path=ROOT / "Resources" / "Ephemeris" / "de440s.bsp",
        iers_path=ROOT / "Resources" / "Ephemeris" / "finals2000A.all",
        leap_seconds_path=ROOT / "Resources" / "Ephemeris" / "Leap_Second.dat",
        allow_predicted_eop=True,
        sample_interval=10.0,
        propagation_model="J2SunMoon",
        propagation_rtol=1e-10,
        propagation_atol_position=1e-4,
        propagation_atol_velocity=1e-7,
    )
    state = np.asarray(current_metadata["inputState"], dtype=float)
    start = parse_absolute_time("2026-09-06 00:00:00", "UTC", "integration")

    generated: dict[str, dict[str, Any]] = {}
    temp_parent = WORK / "temporary_generation"
    temp_parent.mkdir(parents=True, exist_ok=True)
    for hours in (6, 48):
        with tempfile.TemporaryDirectory(prefix=f"{hours}h_", dir=temp_parent) as folder:
            metadata = prepare_environment(
                {
                    "resources": resources,
                    "start_time": start,
                    "duration_seconds": hours * 3600.0,
                    "state_epoch": start,
                    "state_input_mode": "Cartesian",
                    "state_frame": "GCRS",
                    "initial_state": state,
                },
                Path(folder),
            )
            generated[f"{hours}h"] = metadata_summary(metadata)
            table = np.loadtxt(Path(folder) / "environment.txt", comments="#", skiprows=3)
            if hours == 48:
                indices = [int(np.argmin(abs(table[:, 0] - value))) for value in (0.0, 86400.0, 172800.0)]
                day_states = table[indices, 1:4]
                no_repeat = np.linalg.norm(day_states[1] - day_states[0]) > 1000.0 and np.linalg.norm(day_states[2] - day_states[1]) > 1000.0
                check("48 h table does not repeat or hold at 24 h", no_repeat, {"positions_m": day_states.tolist()}, checks)
            check(
                f"generate {hours} h environment",
                metadata["durationSeconds"] == hours * 3600.0
                and metadata["tableCoverageSeconds"][0] < 0
                and metadata["tableCoverageSeconds"][1] > hours * 3600.0,
                generated[f"{hours}h"],
                checks,
            )

    # Same epoch/state, later simulation start: the state must be propagated, not relabelled.
    later_start = add_physical_seconds(start, 600.0)
    with tempfile.TemporaryDirectory(prefix="shifted_epoch_", dir=temp_parent) as folder:
        shifted = prepare_environment(
            {
                "resources": resources,
                "start_time": later_start,
                "duration_seconds": 600.0,
                "state_epoch": start,
                "state_input_mode": "Cartesian",
                "state_frame": "GCRS",
                "initial_state": state,
            },
            Path(folder),
        )
        delta = float(np.linalg.norm(np.asarray(shifted["stateAtSimulationStart"]) - state))
        check(
            "state epoch is propagated to a later simulation start",
            shifted["stateEpochUTC"] != shifted["startUTC"] and delta > 1.0e5,
            {"stateEpochUTC": shifted["stateEpochUTC"], "startUTC": shifted["startUTC"], "stateDifferenceNorm": delta},
            checks,
        )

    # Determinism with the same input and fixed resources.
    hashes = []
    for index in range(2):
        with tempfile.TemporaryDirectory(prefix=f"repeat_{index}_", dir=temp_parent) as folder:
            repeated = prepare_environment(
                {
                    "resources": resources,
                    "start_time": start,
                    "duration_seconds": 600.0,
                    "state_epoch": start,
                    "state_input_mode": "Cartesian",
                    "state_frame": "GCRS",
                    "initial_state": state,
                },
                Path(folder),
            )
            hashes.append(repeated["tableSHA256"])
    check("fixed resources and inputs are deterministic", hashes[0] == hashes[1], {"hashes": hashes}, checks)

    # Missing data and out-of-coverage epochs must fail instead of falling back.
    missing_resources = ResourceSettings(
        spk_path=ROOT / "Resources" / "Ephemeris" / "missing.bsp",
        iers_path=resources.iers_path,
        leap_seconds_path=resources.leap_seconds_path,
        allow_predicted_eop=True,
        sample_interval=10.0,
        propagation_model="J2SunMoon",
        propagation_rtol=1e-10,
        propagation_atol_position=1e-4,
        propagation_atol_velocity=1e-7,
    )
    with tempfile.TemporaryDirectory(prefix="missing_resource_", dir=temp_parent) as folder:
        try:
            prepare_environment(
                {"resources": missing_resources, "start_time": start, "duration_seconds": 60.0,
                 "state_epoch": start, "state_input_mode": "Cartesian", "state_frame": "GCRS", "initial_state": state},
                Path(folder),
            )
            missing_failed = False
        except EphemerisPreparationError:
            missing_failed = True
    check("missing SPK is rejected", missing_failed, {}, checks)

    out_of_range = parse_absolute_time("2099-01-01 00:00:00", "UTC", "integration")
    with tempfile.TemporaryDirectory(prefix="coverage_", dir=temp_parent) as folder:
        try:
            prepare_environment(
                {"resources": resources, "start_time": out_of_range, "duration_seconds": 60.0,
                 "state_epoch": out_of_range, "state_input_mode": "Cartesian", "state_frame": "GCRS", "initial_state": state},
                Path(folder),
            )
            coverage_failed = False
        except EphemerisPreparationError:
            coverage_failed = True
    check("out-of-range IERS epoch is rejected", coverage_failed, {}, checks)

    report = {
        "status": "PASS",
        "checks": checks,
        "generatedEnvironmentVariants": generated,
        "note": "Temporary 6 h/48 h tables were deleted after validation; only summaries and hashes are retained.",
    }
    (WORK / "integration_tests.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    lines = ["# 真实历元与星历环境集成测试", "", f"- 总状态：PASS", f"- 检查数：{len(checks)}", ""]
    lines.extend(f"- `{item['status']}` {item['name']}" for item in checks)
    (WORK / "integration_tests.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(json.dumps({"status": "PASS", "checks": len(checks), "variants": generated}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
