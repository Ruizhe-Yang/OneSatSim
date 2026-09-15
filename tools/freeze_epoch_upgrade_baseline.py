#!/usr/bin/env python3
"""Capture the pre-upgrade Modelica hashes and workbook cell semantics."""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PACKAGE_ROOT = Path(__file__).resolve().parents[1]


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            value.update(block)
    return value.hexdigest()


def serializable(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def workbook_snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells.append({"cell": cell.coordinate, "value": serializable(cell.value), "dataType": cell.data_type})
        payload = json.dumps(cells, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        sheets[sheet.title] = {
            "maxRow": sheet.max_row,
            "maxColumn": sheet.max_column,
            "semanticSHA256": hashlib.sha256(payload.encode("utf-8")).hexdigest(),
            "cells": cells,
        }
    return {"workbook": str(path), "binarySHA256": digest(path), "sheetOrder": workbook.sheetnames, "sheets": sheets}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", default="DesignConfig.xlsx")
    parser.add_argument("--output-dir", default="../work/real_epoch/baseline")
    args = parser.parse_args()
    excel = Path(args.excel)
    if not excel.is_absolute():
        excel = PACKAGE_ROOT / excel
    output_dir = Path(args.output_dir)
    if not output_dir.is_absolute():
        output_dir = (PACKAGE_ROOT / output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    modelica = {
        str(path.relative_to(PACKAGE_ROOT)).replace("\\", "/"): digest(path)
        for path in sorted(PACKAGE_ROOT.rglob("*.mo"))
    }
    (output_dir / "modelica_hashes_before.json").write_text(
        json.dumps(modelica, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (output_dir / "workbook_semantics_before.json").write_text(
        json.dumps(workbook_snapshot(excel), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"Captured {len(modelica)} Modelica files and {len(load_workbook(excel, read_only=True).sheetnames)} workbook sheets")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
