#!/usr/bin/env python3
"""Verify that environment-preparation failure preserves the active generated set."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent / "work" / "real_epoch"
TARGETS = (
    ROOT / "Resources" / "Data" / "GeneratedEphemeris" / "environment.txt",
    ROOT / "Resources" / "Data" / "GeneratedEphemeris" / "environment_metadata.json",
    ROOT / "Scenarios" / "GeneratedScenario.mo",
    ROOT / "Simulation" / "CompleteMission.mo",
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    WORK.mkdir(parents=True, exist_ok=True)
    before = {str(path.relative_to(ROOT)): digest(path) for path in TARGETS}
    with tempfile.TemporaryDirectory(prefix="transaction_", dir=WORK) as folder:
        invalid = Path(folder) / "DesignConfig_invalid_resource.xlsx"
        workbook = load_workbook(ROOT / "DesignConfig.xlsx", data_only=False, read_only=False)
        sheet = workbook["Ephemeris"]
        for row in range(1, sheet.max_row + 1):
            if sheet.cell(row, 1).value == "SPKFile":
                sheet.cell(row, 2).value = "Resources/Ephemeris/definitely_missing.bsp"
                break
        workbook.save(invalid)
        environment = dict(os.environ)
        environment["PYTHONDONTWRITEBYTECODE"] = "1"
        result = subprocess.run(
            [sys.executable, str(ROOT / "tools" / "update_onesatsim_config.py"), "--excel", str(invalid), "--skip-mass-properties"],
            cwd=ROOT,
            env=environment,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
        )
    after = {str(path.relative_to(ROOT)): digest(path) for path in TARGETS}
    passed = result.returncode != 0 and before == after
    report = {
        "status": "PASS" if passed else "FAIL",
        "returnCode": result.returncode,
        "generatedSetPreserved": before == after,
        "hashesBefore": before,
        "hashesAfter": after,
        "updaterOutput": result.stdout,
    }
    (WORK / "transaction_failure_test.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: report[key] for key in ("status", "returnCode", "generatedSetPreserved")}, ensure_ascii=False))
    return 0 if passed else 1


if __name__ == "__main__":
    raise SystemExit(main())
