#!/usr/bin/env python3
"""Compare direct propagated geometry with the generated table at event boundaries."""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path
from typing import Any

import numpy as np
from astropy.utils import iers
from openpyxl import load_workbook
from scipy.interpolate import Akima1DInterpolator

TOOLS = Path(__file__).resolve().parent
ROOT = TOOLS.parent
WORK = ROOT.parent / "work" / "real_epoch"
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from prepare_ephemeris import (  # noqa: E402
    EARTH_RADIUS,
    _dynamics,
    _earth_orientation,
    _make_ephemeris_splines,
    _make_pole_spline,
    _propagate,
    _spk_position,
    add_physical_seconds,
    configure_astropy,
    parse_absolute_time,
)
from jplephem.spk import SPK  # noqa: E402


def crossings(time: np.ndarray, function: np.ndarray, minimum_slope: float) -> tuple[np.ndarray, int]:
    left, right = function[:-1], function[1:]
    indices = np.flatnonzero((left < 0) != (right < 0))
    event_time = time[indices] - left[indices] * (time[indices + 1] - time[indices]) / (right[indices] - left[indices])
    slope = np.abs((right[indices] - left[indices]) / (time[indices + 1] - time[indices]))
    return event_time[slope >= minimum_slope], int(np.count_nonzero(slope < minimum_slope))


def compare_events(label: str, direct: np.ndarray, table: np.ndarray, tolerance: float, results: list[dict[str, Any]]) -> None:
    if len(direct) == 0 and len(table) == 0:
        results.append({"geometry": label, "status": "PASS", "directEvents": 0, "tableEvents": 0, "maxBoundaryError_s": 0.0})
        return
    if len(direct) == 0 or len(table) == 0:
        results.append({"geometry": label, "status": "FAIL", "directEvents": len(direct), "tableEvents": len(table), "maxBoundaryError_s": None})
        return
    nearest = np.array([np.min(np.abs(table - value)) for value in direct])
    maximum = float(np.max(nearest))
    results.append({
        "geometry": label,
        "status": "PASS" if len(direct) == len(table) and maximum <= tolerance else "FAIL",
        "directEvents": int(len(direct)),
        "tableEvents": int(len(table)),
        "maxBoundaryError_s": maximum,
    })


def site_rows(sheet: Any) -> list[dict[str, float | str]]:
    headers = {str(sheet.cell(4, column).value): column for column in range(1, sheet.max_column + 1)}
    rows = []
    for row in range(5, sheet.max_row + 1):
        if sheet.cell(row, headers["Enabled"]).value is True:
            rows.append({
                "name": str(sheet.cell(row, headers["Name"]).value),
                "latitude": math.radians(float(sheet.cell(row, headers["Latitude_deg"]).value)),
                "longitude": math.radians(float(sheet.cell(row, headers["Longitude_deg"]).value)),
                "altitude": float(sheet.cell(row, headers["Altitude_m"]).value or 0.0),
                "minimumElevation": math.radians(float(sheet.cell(row, headers["MinElevation_deg"]).value)),
            })
    return rows


def visibility_function(position: np.ndarray, rotation: np.ndarray, site: dict[str, float | str]) -> np.ndarray:
    latitude = float(site["latitude"])
    longitude = float(site["longitude"])
    fixed = np.array([math.cos(latitude) * math.cos(longitude), math.cos(latitude) * math.sin(longitude), math.sin(latitude)])
    normal = np.einsum("nij,j->ni", rotation, fixed)
    site_position = (EARTH_RADIUS + float(site["altitude"])) * normal
    relative = position - site_position
    elevation_sine = np.einsum("ni,ni->n", relative, normal) / np.linalg.norm(relative, axis=1)
    return elevation_sine - math.sin(float(site["minimumElevation"]))


def main() -> int:
    metadata = json.loads((ROOT / "Resources" / "Data" / "GeneratedEphemeris" / "environment_metadata.json").read_text(encoding="utf-8"))
    configure_astropy(ROOT / "Resources" / "Ephemeris" / "Leap_Second.dat")
    start = parse_absolute_time(metadata["startUTC"], "UTC", "metadata.startUTC")
    duration = float(metadata["durationSeconds"])
    time = np.arange(0.0, duration + 1.0, 1.0)
    margin = float(metadata["requestedSampleInterval"])
    minimum, maximum = -margin, duration + margin
    kernel = SPK.open(str(ROOT / "Resources" / "Ephemeris" / "de440s.bsp"))
    eop = iers.IERS_A.open(str(ROOT / "Resources" / "Ephemeris" / "finals2000A.all"))
    sun_spline, moon_spline = _make_ephemeris_splines(kernel, start, minimum - 300.0, maximum + 300.0)
    pole_spline = _make_pole_spline(start, minimum - 300.0, maximum + 300.0, eop)
    solution = _propagate(
        np.asarray(metadata["inputState"], dtype=float),
        0.0,
        minimum,
        maximum,
        _dynamics(sun_spline, moon_spline, pole_spline),
        1e-10,
        1e-4,
        1e-7,
    )
    direct_position = solution.sol(time)[:3].T
    absolute_times = add_physical_seconds(start, time)
    direct_rotation, _, _, _ = _earth_orientation(absolute_times, eop)
    earth = _spk_position(kernel, "earth", absolute_times)
    direct_sun = _spk_position(kernel, "sun", absolute_times) - earth - direct_position
    direct_sun /= np.linalg.norm(direct_sun, axis=1)[:, None]

    table = np.loadtxt(ROOT / "Resources" / "Data" / "GeneratedEphemeris" / "environment.txt", comments="#", skiprows=3)
    interpolators = [Akima1DInterpolator(table[:, 0], table[:, column]) for column in range(1, 16)]
    table_values = np.column_stack([item(time) for item in interpolators])
    table_position = table_values[:, :3]
    table_sun = table_values[:, 3:6]
    table_sun /= np.linalg.norm(table_sun, axis=1)[:, None]
    table_rotation = table_values[:, 6:15].reshape(len(time), 3, 3)

    results: list[dict[str, Any]] = []
    direct_shadow = np.einsum("ni,ni->n", direct_position, direct_sun) + np.sqrt(np.maximum(0.0, np.sum(direct_position**2, axis=1) - EARTH_RADIUS**2))
    table_shadow = np.einsum("ni,ni->n", table_position, table_sun) + np.sqrt(np.maximum(0.0, np.sum(table_position**2, axis=1) - EARTH_RADIUS**2))
    direct_events, direct_grazing = crossings(time, direct_shadow, 100.0)
    table_events, table_grazing = crossings(time, table_shadow, 100.0)
    compare_events("Eclipse", direct_events, table_events, 2.0, results)

    workbook = load_workbook(ROOT / "DesignConfig.xlsx", data_only=True, read_only=True)
    for kind, sheet_name in (("Ground", "GroundStations"), ("Target", "ImagingTargets")):
        for site in site_rows(workbook[sheet_name]):
            direct_function = visibility_function(direct_position, direct_rotation, site)
            table_function = visibility_function(table_position, table_rotation, site)
            direct_events, _ = crossings(time, direct_function, 1e-6)
            table_events, _ = crossings(time, table_function, 1e-6)
            compare_events(f"{kind}:{site['name']}", direct_events, table_events, 2.0, results)

    kernel.close()
    failed = [item for item in results if item["status"] != "PASS"]
    report = {
        "status": "PASS" if not failed else "FAIL",
        "comparison": "Direct DOP853/JPL/IERS geometry versus final Akima environment table on a 1 s audit grid",
        "nonGrazingBoundaryTolerance_s": 2.0,
        "eclipseGrazingEventsExcluded": direct_grazing + table_grazing,
        "results": results,
        "maximumBoundaryError_s": max((item["maxBoundaryError_s"] or 0.0) for item in results),
    }
    (WORK / "geometry_boundary_validation.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": report["status"], "geometries": len(results), "maximumBoundaryError_s": report["maximumBoundaryError_s"], "failed": failed}, ensure_ascii=False, indent=2))
    return 0 if not failed else 1


if __name__ == "__main__":
    raise SystemExit(main())
