#!/usr/bin/env python3
"""Static checks for the current independent v1.1.4 OneSatSim project."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

import openpyxl

from audit_results import load_contract
from run_provenance import sha256

ROOT = Path(__file__).resolve().parents[1]
REQUIRED_SCRIPTS = (
    "load_package.py", "update_transaction.py", "update_config.py",
    "parameter_schema.py", "parameter_units.py", "prepare_ephemeris.py",
    "generate_mass_properties.py", "run_openmodelica.py", "run_config.py",
    "run_acceptance.py", "check_project.py", "audit_results.py",
    "matv4.py", "generate_dashboard.py", "dashboard_signals.json",
    "generate_embedded_environment.py",
)
GENERATED = (
    "Scenarios/GeneratedSpacecraftDesignConfig.mo",
    "Scenarios/GeneratedScenario.mo",
    "Scenarios/GeneratedEnvironmentData.mo",
    "Scenarios/GeneratedMassProperties.mo",
    "Simulation/CompleteMission.mo",
    "Scripts/ConfigSupport/generated_simulation_flags.txt",
)
PALETTE = {(39, 147, 100), (29, 63, 137), (76, 166, 218),
           (112, 121, 135), (207, 69, 65)}
MSL_CONNECTOR_COLORS = {'real': [0, 0, 127], 'boolean': [255, 0, 255], 'integer': [255, 127, 0], 'electrical': [0, 0, 255], 'thermal': [191, 0, 0], 'rotational': [0, 0, 0], 'frame': [95, 95, 95]}



def _source_digest(root: Path) -> tuple[str, int]:
    files = sorted(root.rglob("*.mo"), key=lambda p: p.relative_to(root).as_posix())
    digest = hashlib.sha256()
    for path in files:
        digest.update(path.relative_to(root).as_posix().encode("utf-8"))
        digest.update(b"\0")
        digest.update(path.read_text(encoding="utf-8-sig").encode("utf-8"))
        digest.update(b"\0")
    return digest.hexdigest(), len(files)


def _order_errors(root: Path) -> list[str]:
    errors = []
    for package in root.rglob("package.mo"):
        folder = package.parent
        order = folder / "package.order"
        if not order.is_file():
            errors.append(f"{order.relative_to(root)} missing")
            continue
        for name in order.read_text(encoding="utf-8").splitlines():
            name = name.strip()
            if name and not (folder / (name + ".mo")).is_file() and not (folder / name / "package.mo").is_file():
                errors.append(f"{order.relative_to(root)}: {name} missing")
    return errors


def _generated_errors(root: Path) -> list[str]:
    path = root / "Scripts/ConfigSupport/generation_manifest.json"
    if not path.is_file():
        return ["generation_manifest.json missing"]
    manifest = json.loads(path.read_text(encoding="utf-8"))
    listed = {
        "DesignConfig.xlsx": manifest.get("workbookSHA256"),
        "Scenarios/DefaultSpacecraftDesignConfig.mo": manifest.get("defaultRecordSHA256"),
        "Scripts/update_config.py": manifest.get("generatorSHA256"),
        "Scripts/parameter_schema.py": manifest.get("schemaSHA256"),
        "Scripts/parameter_units.py": manifest.get("unitRulesSHA256"),
        "Scripts/prepare_ephemeris.py": manifest.get("ephemerisGeneratorSHA256"),
        "Scripts/generate_embedded_environment.py": manifest.get("embeddedGeneratorSHA256"),
        "Scripts/generate_mass_properties.py": manifest.get("massGeneratorSHA256"),
        "Scripts/update_transaction.py": manifest.get("transactionSHA256"),
        **manifest.get("generated", {}),
    }
    errors = [f"{name} hash drift" for name, expected in listed.items()
              if not expected or not (root / name).is_file() or sha256(root / name) != expected]
    errors.extend(f"{name} missing from manifest" for name in GENERATED
                  if name not in manifest.get("generated", {}))
    if not (root / "Scripts/ConfigSupport/current_snapshot.json").is_file():
        errors.append("current_snapshot.json missing")
    return errors


def _parameter_errors(root: Path, parameters: dict) -> list[str]:
    book = openpyxl.load_workbook(root / "DesignConfig.xlsx", read_only=True, data_only=True)
    try:
        if len(book.sheetnames) != 10:
            return [f"Expected ten design sheets, got {len(book.sheetnames)}"]
        rows = {}
        for row_no, row in enumerate(book["Components"].iter_rows(values_only=True), 1):
            if row_no >= 5 and row[0] and row[1] and row[4]:
                rows[str(row[12])] = {"row": row_no,
                                      "value": row[8] if bool(row[9]) else row[6],
                                      "useDefault": bool(row[9])}
    finally:
        book.close()
    catalog = {str(item.get("id")): item for item in parameters.get("componentIds", [])}
    errors = []
    if len(rows) != 454:
        errors.append(f"Expected unchanged 454 component parameter IDs, got {len(rows)}")
    if len(rows) != len(catalog):
        errors.append(f"Workbook IDs {len(rows)} != contract IDs {len(catalog)}")
    for identifier in sorted(set(rows) | set(catalog)):
        workbook = rows.get(identifier)
        item = catalog.get(identifier)
        if workbook is None or item is None:
            errors.append(identifier + " missing")
        elif (item.get("workbook", {}).get("row") != workbook["row"] or
              item.get("selectedWorkbookValue") != workbook["value"] or
              bool(item.get("useDefault")) != workbook["useDefault"]):
            errors.append(identifier + " workbook/contract mismatch")
    return errors


def _dashboard_errors(root: Path, results: dict) -> list[str]:
    catalog = json.loads((root / "Scripts/dashboard_signals.json").read_text(encoding="utf-8"))
    model = (root / "Simulation/OnboardTelemetry.mo").read_text(encoding="utf-8")
    port = (root / "Foundation/Interfaces/OnboardTelemetryPort.mo").read_text(encoding="utf-8")
    saved = {row["path"] for row in results["signals"] if row["save"]}
    expected = set()
    errors = []
    for row in catalog:
        name, source = row["name"], row["source"]
        if not re.search(r"\boutput\s+(?:Real|Integer|Boolean)\s+" + re.escape(name) + r"\b", model):
            errors.append(f"{name} output missing")
        if not re.search(r"\b" + re.escape(name) + r"\s*=\s*telemetry\." + re.escape(source), model):
            errors.append(f"{name} telemetry assignment missing")
        if source.split("[")[0] not in port:
            errors.append(f"{name} telemetry source missing")
        for index in range(1, int(row.get("size", 1)) + 1):
            suffix = f"[{index}]" if row.get("size") else ""
            expected.add(f"onboardTelemetry.{name}{suffix}")
    errors.extend(f"{name} not saved" for name in sorted(expected - saved))
    if not expected:
        errors.append("Dashboard dictionary empty")
    computer_shapes = (
        "Rectangle(extent={{-84,-43},{84,72}}",
        "Rectangle(extent={{-70,-29},{70,58}}",
        "Line(points={{-49,8},{-30,8},{-17,32},{1,-10},{17,17},{49,17}}",
        "Rectangle(extent={{-10,-70},{10,-43}}",
        "Rectangle(extent={{-47,-81},{47,-70}}",
        "Ellipse(extent={{58,-39},{65,-32}}",
    )
    for shape in computer_shapes:
        if model.count(shape) != 2:
            errors.append(f"Computer Icon/Diagram shape missing: {shape}")
    if "DynamicSelect(" in model or "Text(" in model or "KPI" in model:
        errors.append("Old dashboard text remains in computer graphics")
    return errors


def _diagram_errors(root: Path, topology: dict) -> list[str]:
    errors = []
    if topology.get("schema") != "MOSAIC.diagram_topology.v2":
        errors.append("Diagram topology schema mismatch")
    files = {p.stem: p for p in root.rglob("*.mo") if p.stem != "package"}
    if len(list(root.rglob("*.mo"))) != topology.get("expectedMoClassCount"):
        errors.append("Modelica class count differs from v1.1.4 inventory")
    styles = topology.get("connectorStyles", {})
    if len(styles) != 45:
        errors.append("Connector style contract must cover 45 classes")
    for name, spec in {**styles, **topology.get("buStyles", {})}.items():
        path = files.get(name)
        if path is None:
            errors.append(f"{name} missing")
            continue
        source = path.read_text(encoding="utf-8")
        color = "{" + ",".join(str(v) for v in spec["rgb"]) + "}"
        background = ("Rectangle(extent={{-100,-100},{100,100}},"
                      "fillPattern=FillPattern.Solid,fillColor=" + color)
        if source.count(background) != 2:
            errors.append(f"{name} Icon/Diagram background or color mismatch")
        if name in styles and "Text(" in source:
            errors.append(f"{name} connector retains text artwork")

    declaration = re.compile(r"^\s*(?:(?:inner|outer|final|replaceable|parameter|constant|input|output|discrete)\s+)*([A-Za-z_][\w.]*)\s+([A-Za-z_]\w*)(?=[\[\s;(])", re.M)
    members = {}
    for name, path in files.items():
        source = path.read_text(encoding="utf-8")
        split = re.search(r"(?m)^\s*(?:equation|algorithm|initial equation|initial algorithm)\b", source)
        prefix = source[:split.start()] if split else source.split("  annotation(", 1)[0]
        members[name] = {member: typ.split(".")[-1] for typ, member in declaration.findall(prefix)}

    def custom_type(class_name: str, endpoint: str) -> str | None:
        current = class_name
        for token in endpoint.strip().split("."):
            token = re.sub(r"\[.*?\]", "", token)
            current = members.get(current, {}).get(token)
            if current is None:
                return None
        return current if current in styles else None

    connect = re.compile(r"(?m)^\s*connect\(\s*([^,]+?)\s*,\s*([^)]+?)\s*\)\s*annotation\(Line\((.*?)\)\);", re.S)
    all_connect = 0
    custom_connect = 0
    for path in files.values():
        source = path.read_text(encoding="utf-8")
        matches = list(connect.finditer(source))
        if len(matches) != source.count("connect("):
            errors.append(f"{path.relative_to(root)} connect without Line annotation")
        for match in matches:
            all_connect += 1
            ta = custom_type(path.stem, match.group(1))
            tb = custom_type(path.stem, match.group(2))
            line = match.group(3)
            if "smooth=Smooth.None" not in line:
                errors.append(f"{path.relative_to(root)} has non-orthogonal line style")
            if ta and tb and styles[ta]["rgb"] != styles[tb]["rgb"]:
                errors.append(f"{path.relative_to(root)} mismatched custom connector families")
            typ = ta or tb
            if typ:
                custom_connect += 1
                expected = "color={" + ",".join(str(v) for v in styles[typ]["rgb"]) + "}"
                if expected not in line:
                    errors.append(f"{path.relative_to(root)} {match.group(1)} -> {match.group(2)} wrong color")
    if all_connect != topology.get("expectedConnectCount"):
        errors.append(f"connect count {all_connect} differs from 2314 inventory")
    if custom_connect != topology.get("expectedCustomConnectCount"):
        errors.append(f"custom connector line count {custom_connect} differs from inventory")

    space = (root / "Systems/SpacecraftSystem.mo").read_text(encoding="utf-8")
    for name, old in topology.get("topLevelInstances", {}).items():
        found = re.search(r"\b" + re.escape(name) + r"\b[^;]*?transformation\(origin=\{([^}]+)\},\s*extent=\{\{([^}]+)\},\{([^}]+)\}\}", space, re.S)
        if not found:
            errors.append(f"Top level instance {name} absent")
            continue
        location = [float(v) for v in found.group(1).split(",")]
        extent = [[float(v) for v in part.split(",")] for part in found.groups()[1:]]
        if location != old["origin"] or extent != old["extent"]:
            errors.append(f"Top level instance {name} moved or resized")
    if "activeIndexBridge" in space:
        errors.append("SpacecraftSystem retains activeIndexBridge")
    for rel in ("Components/MagnetometerUnit.mo", "Components/StarTrackerYUnit.mo",
                "Components/StarTrackerZUnit.mo", "Components/SunSensorUnit.mo",
                "Simulation/OrbitalEnvironment.mo", "Systems/SpacecraftSystem.mo"):
        source = (root / rel).read_text(encoding="utf-8")
        eq = source.split("\nequation\n", 1)[1].split("\n  annotation(", 1)[0]
        for line in eq.splitlines():
            if line.strip() and not line.strip().startswith("connect("):
                errors.append(f"{rel} retains a non-connect assembly equation")
    for name in ("ASRTUStatePort", "MissionControlBus", "CommandExecutionStatusOutput",
                 "DownlinkPhaseOutput", "MissionModeOutput"):
        if name in files:
            errors.append(f"Unused connector {name} remains")
    for rel, token in (("Systems/N_systems/ElectricalPowerSystem.mo", "CurrentSensor sourceBusCurrent"),
                       ("Components/PowerConditioningUnit.mo", "CurrentSensor pcduHousekeepingCurrent"),
                       ("Components/OnboardComputerUnit.mo", "CurrentSensor currentSensor"),
                       ("Components/HeaterRadiatorUnit.mo", "heaterCurrent[12]")):
        if token in (root / rel).read_text(encoding="utf-8"):
            errors.append(f"{rel} retains removed measurement {token}")
    return errors


def check(root: Path = ROOT) -> dict:
    checks = {}
    def add(name: str, errors: list[str], detail=None) -> None:
        checks[name] = {"status": "FAIL" if errors else "PASS",
                        "detail": detail if detail is not None else errors[:40]}

    package = (root / "package.mo").read_text(encoding="utf-8")
    add("ROOT_IDENTITY", [] if re.search(r"(?m)^package\s+OneSatSim(?:\s|$)", package) and
        'version="1.1.4"' in package and root.name == "OneSatSim" else ["Quoted root/version/directory mismatch"])
    add("FIVE_PACKAGES", [name for name in
        ("Foundation", "Scenarios", "Components", "Systems", "Simulation")
        if not (root / name / "package.mo").is_file()])
    add("PACKAGE_ORDER", _order_errors(root))
    add("SCRIPTS_SINGLE_ROOT", [name for name in REQUIRED_SCRIPTS
        if not (root / "Scripts" / name).is_file()] +
        (["Legacy tools/ tree present"] if (root / "tools").exists() else []))
    add("FLAGS_SINGLE_SOURCE", [] if (root / "Scripts/ConfigSupport/generated_simulation_flags.txt").is_file()
        and not (root / "Simulation/generated_simulation_flags.txt").exists()
        else ["Generated simulation flags absent or duplicated"])
    runtime_errors = ["Resources/ directory remains"] if (root / "Resources").exists() else []
    for source in root.rglob("*.mo"):
        content = source.read_text(encoding="utf-8")
        if any(term in content for term in ("tableOnFile=true", "loadResource(", "environmentFileOverride", "modelica://")):
            runtime_errors.append(f"{source.relative_to(root)} retains an external table reference")
    add("PURE_MO", runtime_errors)
    add("GENERATED_GROUP", _generated_errors(root))
    contracts = {}
    errors = []
    for name in ("parameters.json", "results.json", "diagram_topology.json"):
        path = root / "Scripts/Contracts" / name
        try:
            contracts[name] = load_contract(path) if name == "results.json" else json.loads(
                path.read_text(encoding="utf-8"))
            if not (contracts[name].get("schemaVersion") or
                    (name == "diagram_topology.json" and contracts[name].get("schema"))):
                errors.append(name + " schema identity missing")
        except Exception as exc:
            errors.append(f"{name}: {exc}")
    add("CONTRACT_FILES", errors, {"errors": errors, "files": list(contracts)})
    if "parameters.json" in contracts:
        add("PARAMETER_CONTRACT", _parameter_errors(root, contracts["parameters.json"]))
    else:
        add("PARAMETER_CONTRACT", ["Parameters contract unavailable"])
    top = (root / "Simulation/CompleteMission.mo").read_text(encoding="utf-8")
    expected = ("OrbitalEnvironment orbitalEnvironment", "Systems.SpacecraftSystem spacecraft",
                "OnboardTelemetry onboardTelemetry")
    top_errors = [name + " absent" for name in expected if name not in top]
    if len(re.findall(r"\bconnect\s*\(", top)) != 2:
        top_errors.append("CompleteMission must have exactly two connects")
    if any(name in top for name in ("TaskRequest", "RuntimeDisplay", "ParameterDisplay")):
        top_errors.append("Legacy top-level runtime object present")
    add("TOP_LEVEL", top_errors)
    modelica = [path for path in root.rglob("*.mo")]
    stale = [str(path.relative_to(root)) for path in modelica
             if "allowNewRequests" in path.read_text(encoding="utf-8") or
             "'MOSAIC-12UCubeSat'" in path.read_text(encoding="utf-8")]
    stale += [name for name in ("TaskRequest", "RuntimeDisplay", "ParameterDisplay")
              if (root / "Simulation" / (name + ".mo")).exists()]
    add("NO_EXTERNAL_REQUEST", stale)
    env = (root / "Foundation/Interfaces/EnvironmentPort.mo").read_text(encoding="utf-8")
    space = (root / "Systems/SpacecraftSystem.mo").read_text(encoding="utf-8")
    orbit = (root / "Simulation/OrbitalEnvironment.mo").read_text(encoding="utf-8")
    env_errors = []
    if "connect(orbitalEnvironment.environment,spacecraft.environment)" not in top:
        env_errors.append("Root EnvironmentPort connection missing")
    if "EnvironmentPort environment" not in space or "EnvironmentPort environment" not in orbit:
        env_errors.append("Spacecraft/environment common port absent")
    for member in ("activeTargetIndex", "activeGroundStationIndex"):
        if member not in env:
            env_errors.append(member + " absent")
    add("ENVIRONMENT_PORT", env_errors)
    if "results.json" in contracts:
        add("DASHBOARD_SAVED", _dashboard_errors(root, contracts["results.json"]))
    else:
        add("DASHBOARD_SAVED", ["Results contract unavailable"])
    if "diagram_topology.json" in contracts:
        add("DIAGRAM_TOPOLOGY", _diagram_errors(root, contracts["diagram_topology.json"]))
    else:
        add("DIAGRAM_TOPOLOGY", ["Diagram topology contract unavailable"])
    return {"schemaVersion": "1.1.0", "project": str(root), "checks": checks,
            "status": "PASS" if all(item["status"] == "PASS" for item in checks.values()) else "FAIL"}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    report = check()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"status": report["status"],
                      "checks": {key: value["status"] for key, value in report["checks"].items()}},
                     ensure_ascii=False))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
